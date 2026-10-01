"""Plain cover letter output: a simple black-and-white .docx and matching PDF.

Layout is a standard business letter (name + contact, date, company, body),
Times New Roman 11pt, 1-inch margins, no color. The .docx is there so the
letter can be opened and tweaked in Word; the PDF is what gets uploaded.
"""

from __future__ import annotations

import html
import re
from datetime import date
from pathlib import Path

FONT = "Times New Roman"
FONT_PT = 11


def _format_phone(phone: str) -> str:
    digits = "".join(ch for ch in phone if ch.isdigit())
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}" if len(digits) == 10 else phone


def _header_lines(profile: dict) -> tuple[str, str]:
    p = profile.get("personal", {})
    contact = [p.get("email", ""), _format_phone(p.get("phone", "")), p.get("linkedin_url", ""),
               p.get("github_url", "")]
    contact = [c.replace("https://", "").replace("www.", "").rstrip("/") for c in contact if c]
    return p.get("full_name", ""), "  |  ".join(contact)


def _segments(line: str) -> list[tuple[str, bool]]:
    """Split a line into (text, is_bold) runs using **bold** markers."""
    parts = re.split(r"(\*\*[^*]+\*\*)", line)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]


def _paragraphs(letter: str) -> list[list[str]]:
    """Blank-line separated paragraphs, each a list of lines (keeps 'Sincerely,' / name apart)."""
    return [[ln.strip() for ln in block.splitlines() if ln.strip()]
            for block in letter.strip().split("\n\n") if block.strip()]


def write_docx(letter: str, profile: dict, company: str, out: Path) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(FONT_PT)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)

    name, contact = _header_lines(profile)
    run = doc.add_paragraph().add_run(name)
    run.bold = True
    run.font.size = Pt(13)
    doc.add_paragraph(contact)
    doc.add_paragraph(date.today().strftime("%B %d, %Y").replace(" 0", " "))
    if company:
        doc.add_paragraph(company)
    for lines in _paragraphs(letter):
        para = doc.add_paragraph()
        for i, line in enumerate(lines):
            for text, bold in _segments(line):
                run = para.add_run(text)
                run.bold = bold
            if i < len(lines) - 1:
                run.add_break()
    doc.save(str(out))
    return out


def write_pdf(letter: str, profile: dict, company: str, out: Path) -> Path:
    from applypilot.scoring.pdf import render_pdf

    name, contact = _header_lines(profile)
    def line_html(line: str) -> str:
        return "".join(f"<b>{html.escape(t)}</b>" if b else html.escape(t) for t, b in _segments(line))

    body = "".join("<p>" + "<br>".join(line_html(ln) for ln in lines) + "</p>"
                   for lines in _paragraphs(letter))
    today = date.today().strftime("%B %d, %Y").replace(" 0", " ")
    page = f"""<!doctype html><html><head><meta charset="utf-8"><style>
      @page {{ size: Letter; }}
      body {{ font-family: "{FONT}", "Liberation Serif", "DejaVu Serif", serif;
             font-size: {FONT_PT}pt; line-height: 1.35; color: #000; margin: 0; }}
      p {{ margin: 0 0 10pt 0; }}
      .name {{ font-weight: bold; font-size: 13pt; margin-bottom: 2pt; }}
    </style></head><body>
      <p class="name">{html.escape(name)}</p>
      <p>{html.escape(contact)}</p>
      <p>{today}</p>
      {f"<p>{html.escape(company)}</p>" if company else ""}
      {body}
    </body></html>"""
    render_pdf(page, str(out), margin="1in")
    return out


def render_letter(txt_path: Path, profile: dict, company: str = "") -> Path:
    """Write <name>.docx and <name>.pdf next to the letter's .txt. Returns the PDF path."""
    txt_path = Path(txt_path)
    letter = txt_path.read_text(encoding="utf-8")
    write_docx(letter, profile, company, txt_path.with_suffix(".docx"))
    return write_pdf(letter, profile, company, txt_path.with_suffix(".pdf"))
